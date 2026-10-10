"""Atlanta Fed Market Probability Tracker (DS-88) and Cleveland Fed inflation nowcasts (DS-90).

Neither publishes a documented API. Both adapters save the downloaded file to raw/
first, then parse it by header text rather than fixed positions, so a layout change
fails loudly (and can be re-parsed from raw/) instead of storing wrong numbers.

- Atlanta Fed: the page links an "MPT Historical Data" workbook. Its URL goes in the
  ATLANTA_MPT_URL environment variable once confirmed (set by Terraform variable
  atlanta_mpt_url). The parser finds the header row (the first row with a "date"
  column) and returns long rows (date, column, value).
- Cleveland Fed: the nowcast page has one table per basis. Rows are months or quarters.
  Columns are CPI, core CPI, PCE and core PCE. Captions separate month-over-month,
  year-over-year and annualized quarterly percent changes. A blank cell is a released
  print and is skipped. The current forecast is not stored as a historical consensus.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from html.parser import HTMLParser

ATLANTA_PAGE = "https://www.atlantafed.org/cenfis/market-probability-tracker"
CLEVELAND_PAGE = "https://www.clevelandfed.org/indicators-and-data/inflation-nowcasting"
MEASURES = ("Core CPI", "CPI", "Core PCE", "PCE")  # longest match first


class LayoutChangedError(ValueError):
    """The source file no longer looks like what the parser expects."""


def parse_atlanta_xlsx(content: bytes, sheet: str | None = None) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    header, rows = None, []
    for values in ws.iter_rows(values_only=True):
        cells = ["" if v is None else v for v in values]
        if header is None:
            if any(str(c).strip().lower() == "date" for c in cells):
                header = [str(c).strip() for c in cells]
            continue
        rec = dict(zip(header, cells, strict=False))
        d = rec.get(next(h for h in header if h.lower() == "date"))
        if isinstance(d, datetime):
            d = d.date()
        elif isinstance(d, str) and d:
            d = date.fromisoformat(d[:10])
        if not isinstance(d, date):
            continue
        for col, val in rec.items():
            if col and col.lower() != "date" and isinstance(val, (int, float)):
                rows.append({"date": d, "column": col, "value": float(val), "source_id": "DS-88"})
    if header is None:
        raise LayoutChangedError("Atlanta Fed MPT workbook has no header row with a 'Date' column")
    return rows


class _Tables(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tables: list[dict] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None
        self._in_caption = False

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.tables.append({"caption": "", "rows": []})
            self._row = None
            self._cell = None
            self._in_caption = False
        elif tag == "caption" and self.tables:
            self._in_caption = True
        elif tag == "tr" and self.tables:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag):
        if tag == "caption":
            self._in_caption = False
        elif tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None and self.tables:
            if any(self._row):
                self.tables[-1]["rows"].append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._in_caption and self.tables:
            self.tables[-1]["caption"] += data
        elif self._cell is not None:
            self._cell.append(data)


def _basis(caption: str) -> str | None:
    text = " ".join(caption.lower().split())
    if "year-over-year" in text or "year over year" in text:
        return "yoy"
    if "annualized" in text:
        return "annualized_quarterly"
    if "month-over-month" in text or "month over month" in text:
        return "mom"
    return None


def _percent(cell: str) -> float | None:
    text = cell.strip().replace("%", "")
    if not text or text.lower() in {"n/a", "na", "-"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _measure(label: str) -> str | None:
    for m in MEASURES:
        if label.strip().lower().startswith(m.lower()):
            return m
    return None


def parse_cleveland_html(html: str, as_of: date) -> list[dict]:
    """Rows: as_of, basis (mom, yoy, annualized_quarterly), period, measure, value (percent).

    Column headers are the measures. The row label is the month or quarter. Updated
    and footnote cells are ignored. A blank cell is a released print, not a zero.
    """
    p = _Tables()
    p.feed(html)
    rows = []
    for table in p.tables:
        basis = _basis(table["caption"])
        body = table["rows"]
        if not basis or len(body) < 2:
            continue
        header = body[0]
        for record in body[1:]:
            if not record:
                continue
            period = record[0]
            if not period or period.lower().startswith("note"):
                continue
            for name, cell in zip(header[1:], record[1:], strict=False):
                measure = _measure(name)
                value = _percent(cell)
                if not measure or value is None:
                    continue
                rows.append(
                    {
                        "as_of": as_of,
                        "basis": basis,
                        "period": period,
                        "measure": measure,
                        "value": value,
                        "source_id": "DS-90",
                    }
                )
    if not rows:
        raise LayoutChangedError("No CPI/PCE nowcast rows found in the Cleveland Fed page")
    return rows
