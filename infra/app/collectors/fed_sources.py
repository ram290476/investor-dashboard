"""Atlanta Fed Market Probability Tracker (DS-88) and Cleveland Fed inflation nowcasts (DS-90).

Neither publishes a documented API. Both adapters save the downloaded file to raw/
first, then parse it by header text rather than fixed positions, so a layout change
fails loudly (and can be re-parsed from raw/) instead of storing wrong numbers.

- Atlanta Fed: the page links an "MPT Historical Data" workbook. Its URL goes in the
  ATLANTA_MPT_URL environment variable once confirmed (set by Terraform variable
  atlanta_mpt_url). The parser finds the header row (the first row with a "date"
  column) and returns long rows (date, column, value).
- Cleveland Fed: the nowcast page shows tables for CPI, core CPI, PCE and core PCE.
  The parser reads every HTML table and keeps rows whose first cell names one of those
  measures, with the column headers as periods.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from html.parser import HTMLParser

ATLANTA_PAGE = "https://www.atlantafed.org/cenfis/market-probability-tracker"
CLEVELAND_PAGE = "https://www.clevelandfed.org/indicators-and-data/inflation-nowcasting"
MEASURES = ("Core CPI", "CPI", "Core PCE", "PCE")  # longest match first


class LayoutChangedError(ValueError):
    """The source file no longer looks like what the parser expects."""


def parse_atlanta_xlsx(content: bytes, sheet: str | None = None) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    header, rows = None, []
    for values in ws.iter_rows(values_only=True):
        cells = ["" if v is None else v for v in values]
        if header is None:
            if any(str(c).strip().lower() == "date" for c in cells):
                header = [str(c).strip() for c in cells]
            continue
        rec = dict(zip(header, cells, strict=False))
        d = rec.get(next(h for h in header if h.lower() == "date"))
        if isinstance(d, datetime):
            d = d.date()
        elif isinstance(d, str) and d:
            d = date.fromisoformat(d[:10])
        if not isinstance(d, date):
            continue
        for col, val in rec.items():
            if col and col.lower() != "date" and isinstance(val, (int, float)):
                rows.append({"date": d, "column": col, "value": float(val), "source_id": "DS-88"})
    if header is None:
        raise LayoutChangedError("Atlanta Fed MPT workbook has no header row with a 'Date' column")
    return rows


class _Tables(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tables: list[list[list[str]]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.tables.append([])
        elif tag == "tr" and self.tables:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None and self.tables:
            self.tables[-1].append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def _measure(label: str) -> str | None:
    for m in MEASURES:
        if label.strip().lower().startswith(m.lower()):
            return m
    return None


def parse_cleveland_html(html: str, as_of: date) -> list[dict]:
    """Rows: as_of, table (index), period (column header), measure, value (percent)."""
    p = _Tables()
    p.feed(html)
    rows = []
    for t_idx, table in enumerate(p.tables):
        if not table:
            continue
        header = table[0]
        for r in table[1:]:
            if not r:
                continue
            measure = _measure(r[0])
            if not measure:
                continue
            for period, cell in zip(header[1:], r[1:], strict=False):
                try:
                    val = float(cell.replace("%", ""))
                except ValueError:
                    continue
                rows.append(
                    {
                        "as_of": as_of,
                        "table": t_idx,
                        "period": period,
                        "measure": measure,
                        "value": val,
                        "source_id": "DS-90",
                    }
                )
    if not rows:
        raise LayoutChangedError("No CPI/PCE nowcast rows found in the Cleveland Fed page")
    return rows
